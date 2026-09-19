"""Read-only extraction of the supplied workbook into a regression fixture."""
import json
import sys
from pathlib import Path
import openpyxl

source = Path(sys.argv[1])
book = openpyxl.load_workbook(source, data_only=True)
formulas = openpyxl.load_workbook(source, data_only=False)
sheet, answers = book['Решение'], book['Ответы']
metrics = ['balance', 'range', 'fear', 'uncertainty', 'calm', 'interest', 'result']
fixture = {
    'source': source.name,
    'options': [sheet.cell(5, c).value for c in [7, 9, 11, 13]],
    'pairs': [{
        'negativeLabel': sheet.cell(r, 3).value,
        'positiveLabel': sheet.cell(r+1, 3).value,
        'weight': sheet.cell(r, 4).value,
        'negative': [sheet.cell(r, c).value for c in [7, 9, 11, 13]],
        'positive': [sheet.cell(r+1, c).value for c in [7, 9, 11, 13]],
    } for r in range(7, 37, 2)],
    'expected': [{metric: answers.cell(r, c).value for r, metric in enumerate(metrics, 6)}
                 for c in range(3, 7)],
    'formulas': {metric: formulas['Ответы'].cell(r, 3).value for r, metric in enumerate(metrics, 6)},
}
destination = Path(__file__).resolve().parents[1] / 'tests/fixtures/excel.json'
destination.parent.mkdir(parents=True, exist_ok=True)
destination.write_text(json.dumps(fixture, ensure_ascii=False, indent=2) + '\n')
print(f'Extracted {len(fixture["pairs"])} pairs and all 28 expected metrics.')
